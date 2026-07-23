from __future__ import annotations

import json
import re
import zipfile
from pathlib import Path

from docx import Document

from v13_content import SCENES, TARGET_SECONDS, validate


ROOT = Path(__file__).resolve().parents[2]
DOCX = ROOT / "文本" / "奔月_正式分镜稿_V1.3.docx"
OUTPUT = Path(__file__).with_name("v13_independent_verification.json")


def collect_text(doc: Document) -> str:
    items = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            items.extend(cell.text for cell in row.cells)
    return "\n".join(items)


def main():
    total, source_shots = validate()
    doc = Document(str(DOCX))
    all_text = collect_text(doc)

    scene_headings = [
        p.text for p in doc.paragraphs
        if re.match(r"^第\d+场\s+", p.text)
    ]

    shot_ids = []
    empty_cells = []
    for table_index, table in enumerate(doc.tables):
        for row_index, row in enumerate(table.rows):
            values = [cell.text.strip() for cell in row.cells]
            if values and re.fullmatch(r"\d{3}", values[0]):
                shot_ids.append(values[0])
                if any(not value for value in values):
                    empty_cells.append([table_index, row_index, values[0]])

    with zipfile.ZipFile(DOCX) as package:
        header_footer = {
            name: package.read(name).decode("utf-8", errors="replace")
            for name in package.namelist()
            if name.startswith("word/header") or name.startswith("word/footer")
        }

    expected_shots = [f"{i:03d}" for i in range(1, source_shots + 1)]
    checks = {
        "docx_exists": DOCX.exists(),
        "bytes": DOCX.stat().st_size,
        "source_seconds": total,
        "target_seconds": TARGET_SECONDS,
        "source_scene_count": len(SCENES),
        "docx_scene_heading_count": len(scene_headings),
        "source_shot_count": source_shots,
        "docx_shot_count": len(shot_ids),
        "shot_ids_contiguous": shot_ids == expected_shots,
        "shot_ids_unique": len(shot_ids) == len(set(shot_ids)),
        "empty_shot_cells": empty_cells,
        "table_count": len(doc.tables),
        "section_count": len(doc.sections),
        "metadata_title": doc.core_properties.title,
        "required_terms": {
            "三段直播": "三段直播" in all_text,
            "2029年4月6日": "2029年4月6日" in all_text,
            "2029年4月8日": "2029年4月8日" in all_text,
            "2029年4月13日": "2029年4月13日" in all_text,
            "东风着陆场": "东风着陆场" in all_text,
            "鹊桥二号": "鹊桥二号" in all_text,
            "环月大椭圆冻结轨道": "环月大椭圆冻结轨道" in all_text,
            "天链": "天链" in all_text,
            "着陆 T+06:00": "着陆 T+06:00" in all_text,
            "T+28小时": "T+28小时" in all_text,
            "无配乐": "无配乐" in all_text,
        },
        "forbidden_active_wording_absent": {
            "细小蓝点": "细小蓝点" not in all_text,
            "望月呼号": "望月" not in all_text,
            "青岛VLBI核心": "青岛VLBI核心" not in all_text,
            "服务舱观察室": "服务舱观察室" not in all_text,
        },
        "header_footer_v13": all("V1.3" in xml for xml in header_footer.values()),
        "header_footer_parts": sorted(header_footer),
    }

    checks["passed"] = all([
        checks["source_seconds"] == checks["target_seconds"] == 840,
        checks["source_scene_count"] == checks["docx_scene_heading_count"] == 15,
        checks["source_shot_count"] == checks["docx_shot_count"] == 100,
        checks["shot_ids_contiguous"],
        checks["shot_ids_unique"],
        not checks["empty_shot_cells"],
        all(checks["required_terms"].values()),
        all(checks["forbidden_active_wording_absent"].values()),
        checks["header_footer_v13"],
    ])

    OUTPUT.write_text(json.dumps(checks, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(checks, ensure_ascii=False, indent=2))
    if not checks["passed"]:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
