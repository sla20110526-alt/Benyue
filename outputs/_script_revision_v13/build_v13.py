from __future__ import annotations

import json
import sys
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

from v13_content import (
    ARCHITECTURE,
    CREW,
    GROUND,
    LANDING_DATE,
    LANDING_SITE,
    MEDIA,
    PROFESSIONAL_CHECKS,
    RETURN_SITE,
    SCENES,
    SOURCES,
    SUBTITLE,
    TARGET_SECONDS,
    TITLE,
    VERSION,
    validate,
)


ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "文本" / "奔月_正式分镜稿_V1.2.docx"
OUTPUT = ROOT / "文本" / "奔月_正式分镜稿_V1.3.docx"
QA_OUTPUT = Path(__file__).with_name("v13_qa_report.json")

FONT = "Noto Sans CJK SC"
BLUE = "17365D"
MID_BLUE = "3F6E9E"
LIGHT_BLUE = "DCEAF7"
LIGHTER_BLUE = "EDF4FA"
GRAY = "E7E9ED"
LIGHT_GRAY = "F5F6F8"
WHITE = "FFFFFF"
RED = "F4CCCC"
YELLOW = "FFF2CC"
GREEN = "D9EAD3"


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=55, start=70, bottom=55, end=70):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def keep_row_together(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_repeat_table_header(row):
    repeat_table_header(row)
    keep_row_together(row)


def set_cell_width(cell, width_in: float):
    cell.width = Inches(width_in)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_in * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def set_run_font(run, size: float, bold=False, color="1A1A1A", italic=False):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def clear_body(doc: Document):
    body = doc._element.body
    sect_pr = body.sectPr
    for child in list(body):
        if child is not sect_pr:
            body.remove(child)


def clear_paragraph_content(paragraph):
    p_pr = paragraph._p.pPr
    for child in list(paragraph._p):
        if child is not p_pr:
            paragraph._p.remove(child)


def configure_document(doc: Document):
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11.69)
    section.page_height = Inches(8.27)
    section.top_margin = Inches(0.45)
    section.bottom_margin = Inches(0.45)
    section.left_margin = Inches(0.45)
    section.right_margin = Inches(0.45)
    section.header_distance = Inches(0.18)
    section.footer_distance = Inches(0.18)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(8.5)
    pf = normal.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(2.5)
    pf.line_spacing = 1.02

    if "Scene Heading" in [s.name for s in doc.styles]:
        heading = doc.styles["Scene Heading"]
    else:
        heading = doc.styles.add_style("Scene Heading", 1)
    heading.font.name = FONT
    heading._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    heading.font.size = Pt(12.5)
    heading.font.bold = True
    heading.font.color.rgb = RGBColor.from_string(BLUE)
    heading.paragraph_format.space_before = Pt(5)
    heading.paragraph_format.space_after = Pt(4)
    heading.paragraph_format.keep_with_next = True

    header = section.header
    hp = header.paragraphs[0]
    clear_paragraph_content(hp)
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hr = hp.add_run(f"《奔月》正式分镜稿 {VERSION}")
    set_run_font(hr, 7, color="6B7280")

    footer = section.footer
    p = footer.paragraphs[0]
    clear_paragraph_content(p)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(f"《奔月》正式分镜稿 {VERSION}  ｜  第 ")
    set_run_font(r, 7, color="6B7280")
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    p._p.append(fld)
    r2 = p.add_run(" 页")
    set_run_font(r2, 7, color="6B7280")


def add_paragraph(doc, text="", size=8.5, bold=False, color="1A1A1A", align=None, italic=False,
                  before=0, after=2.5, keep=False, style=None):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.keep_with_next = keep
    p.paragraph_format.line_spacing = 1.02
    if text:
        run = p.add_run(text)
        set_run_font(run, size, bold=bold, color=color, italic=italic)
    return p


def add_heading(doc, text):
    p = doc.add_paragraph(style="Scene Heading")
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.page_break_before = False
    run = p.add_run(text)
    set_run_font(run, 12.5, bold=True, color=BLUE)
    return p


def style_table(table, widths, header=True, font_size=7.8):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for ridx, row in enumerate(table.rows):
        keep_row_together(row)
        if header and ridx == 0:
            set_repeat_table_header(row)
        for cidx, cell in enumerate(row.cells):
            set_cell_width(cell, widths[cidx])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if header and ridx == 0:
                set_cell_shading(cell, BLUE)
            elif ridx % 2 == 0:
                set_cell_shading(cell, LIGHT_GRAY)
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                for run in p.runs:
                    set_run_font(run, font_size, bold=(header and ridx == 0), color=(WHITE if header and ridx == 0 else "1A1A1A"))


def set_cell_text(cell, text, size=7.8, bold=False, color="1A1A1A", align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run(str(text))
    set_run_font(run, size, bold=bold, color=color)


def add_kv_table(doc, rows, widths=(1.55, 9.24)):
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for idx, (key, value) in enumerate(rows):
        cells = table.add_row().cells
        set_cell_text(cells[0], key, 8, bold=True, color=BLUE)
        set_cell_text(cells[1], value, 8)
        set_cell_shading(cells[0], LIGHT_BLUE)
        set_cell_shading(cells[1], WHITE if idx % 2 == 0 else LIGHT_GRAY)
        for i, cell in enumerate(cells):
            set_cell_width(cell, widths[i])
            set_cell_margins(cell, top=60, bottom=60)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        keep_row_together(table.rows[-1])
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def timecode(seconds: int) -> str:
    return f"{seconds // 60:02d}:{seconds % 60:02d}"


def add_cover(doc, total_shots):
    add_paragraph(doc, TITLE, size=26, bold=True, color=BLUE, align=WD_ALIGN_PARAGRAPH.CENTER, before=8, after=2)
    add_paragraph(doc, SUBTITLE, size=12, color=MID_BLUE, align=WD_ALIGN_PARAGRAPH.CENTER, after=4)
    add_paragraph(doc, f"{VERSION}  |  正片 14分00秒  |  {len(SCENES)}场  |  {total_shots}镜", size=10.5, bold=True,
                  color="26384A", align=WD_ALIGN_PARAGRAPH.CENTER, after=6)
    add_paragraph(
        doc,
        "说明：本片为中国首次载人登月未来纪实推演。人物、日期、呼号、口令及未公开任务细节均为艺术化创作；航天器构型、任务链路、月面物理与返回流程以截至2026年7月公开资料为边界。片尾说明不计入14分钟正片。",
        size=8.5,
        color="4B5563",
        align=WD_ALIGN_PARAGRAPH.CENTER,
        after=5,
    )
    add_kv_table(doc, [
        ("正式基线", "文本/奔月_正式分镜稿_V1.3.docx；V1.2及以前版本转为历史基线。"),
        ("任务日期", LANDING_DATE),
        ("着陆地点", LANDING_SITE),
        ("返回地点", RETURN_SITE),
        ("画面规格", "16:9；成片3840×2160；24fps；未来纪实、真实工程质感、AI真人。"),
        ("声音锁定", "全片无配乐；对白、无线电、设备声与任务通信随镜头同步生成；月面外部无空气传播声。"),
        ("直播结构", "三段直播：月轨对接；落月—月面任务—再对接；离月返回—再入—东风回收。长时间段用字幕、任务回放和时间压缩跨越。"),
        ("艺术光照", "保留2029年4月8日与锁定坐标，沿用已确认的艺术化月面硬日光，并在片头片尾推演声明中承担偏差。"),
    ])


def add_control_section(doc):
    add_heading(doc, "一、全片制作总控")
    add_kv_table(doc, [
        ("核心叙事", "三人乘组在完整国家系统支撑下完成月轨转移、两人登月、三人重聚和平安返航。"),
        ("张力来源", "不设置重大险情；张力来自对接、分离、落点选择、触月、起飞、再对接、再入黑障和搜救等不可逆节点。"),
        ("事件密度", "每30—60秒至少出现1—2个可识别事件点：状态变化、选择、授权、时间跃进或系统闭合。"),
        ("时间压缩", "48小时准备、着陆后6小时整备、约28小时月面驻留、数小时再交会和三天地月巡航使用明确字幕、任务回放和匹配剪辑；关键节点保持实时质感。"),
        ("AI边界", "小天负责传感器融合、走廊预测、候选区排序和异常提示；最终选择、授权和中止权归乘组与北京。"),
        ("声学边界", "全片不配乐；太空和月面外部不使用空气传播的发动机声、脚步声、尘土声、车轮声或旗帜飘动声。"),
        ("文字策略", "任务日期、地点、长时间跃进和片尾推演声明由后期字幕完成；AI原始画面不生成可读文字。"),
        ("直播解释", "记者和三名专业方向专家负责事件前解释、时间桥接与事件后复盘；对接、最终下降、第一步、起飞、再对接、黑障恢复和着陆由任务通信主导。"),
    ])


def add_roles_section(doc):
    add_heading(doc, "二、航天器构型、人员岗位与地面席位")
    add_paragraph(doc, "2.1 航天器构型", size=9, bold=True, color=BLUE, keep=True)
    table = doc.add_table(rows=1, cols=2)
    headers = ["系统／舱段", "功能与叙事边界"]
    for i, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], h, 8, bold=True, color=WHITE, align=WD_ALIGN_PARAGRAPH.CENTER)
    for name, desc in ARCHITECTURE:
        cells = table.add_row().cells
        set_cell_text(cells[0], name, 7.8, bold=True, color=BLUE)
        set_cell_text(cells[1], desc, 7.8)
    style_table(table, [1.75, 9.04], header=True, font_size=7.8)

    add_paragraph(doc, "2.2 乘组与任务角色", size=9, bold=True, color=BLUE, keep=True, before=3)
    table = doc.add_table(rows=1, cols=3)
    for i, h in enumerate(["人物", "岗位", "职责"]):
        set_cell_text(table.rows[0].cells[i], h, 8, bold=True, color=WHITE, align=WD_ALIGN_PARAGRAPH.CENTER)
    for name, role, desc in CREW:
        cells = table.add_row().cells
        set_cell_text(cells[0], name, 7.8, bold=True, color=BLUE)
        set_cell_text(cells[1], role, 7.8)
        set_cell_text(cells[2], desc, 7.8)
    style_table(table, [1.0, 2.2, 7.59], header=True, font_size=7.8)

    add_paragraph(doc, "2.3 直播人物", size=9, bold=True, color=BLUE, keep=True, before=3)
    table = doc.add_table(rows=1, cols=3)
    for i, h in enumerate(["人物", "岗位", "职责"]):
        set_cell_text(table.rows[0].cells[i], h, 8, bold=True, color=WHITE, align=WD_ALIGN_PARAGRAPH.CENTER)
    for name, role, desc in MEDIA:
        cells = table.add_row().cells
        set_cell_text(cells[0], name, 7.8, bold=True, color=BLUE)
        set_cell_text(cells[1], role, 7.8)
        set_cell_text(cells[2], desc, 7.8)
    style_table(table, [1.0, 2.2, 7.59], header=True, font_size=7.8)

    add_paragraph(doc, "2.4 地面系统与艺术化呼号", size=9, bold=True, color=BLUE, keep=True, before=3)
    table = doc.add_table(rows=1, cols=3)
    for i, h in enumerate(["呼号", "系统", "片中作用"]):
        set_cell_text(table.rows[0].cells[i], h, 8, bold=True, color=WHITE, align=WD_ALIGN_PARAGRAPH.CENTER)
    for call, system, role in GROUND:
        cells = table.add_row().cells
        set_cell_text(cells[0], call, 7.8, bold=True, color=BLUE)
        set_cell_text(cells[1], system, 7.8)
        set_cell_text(cells[2], role, 7.8)
    style_table(table, [1.15, 2.55, 7.09], header=True, font_size=7.8)


def add_overview(doc):
    add_heading(doc, "三、场次与时长总览")
    table = doc.add_table(rows=1, cols=6)
    headers = ["场次", "场名", "时间码", "时长", "地点", "主要事件点"]
    for i, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], h, 8, bold=True, color=WHITE, align=WD_ALIGN_PARAGRAPH.CENTER)
    cursor = 0
    for idx, scene in enumerate(SCENES, start=1):
        cells = table.add_row().cells
        values = [idx, scene["title"], f"{timecode(cursor)}–{timecode(cursor + scene['duration'])}", f"{scene['duration']}秒",
                  scene["location"], scene["beat"]]
        for cidx, value in enumerate(values):
            set_cell_text(cells[cidx], value, 7.4, bold=(cidx == 1), color=(BLUE if cidx == 1 else "1A1A1A"),
                          align=WD_ALIGN_PARAGRAPH.CENTER if cidx in (0, 2, 3) else WD_ALIGN_PARAGRAPH.LEFT)
        cursor += scene["duration"]
    style_table(table, [0.48, 1.18, 1.05, 0.55, 2.15, 5.38], header=True, font_size=7.4)
    add_paragraph(doc, f"总计：{len(SCENES)}场，{sum(len(scene['shots']) for scene in SCENES)}镜，{TARGET_SECONDS}秒（14分00秒）。",
                  size=8.5, bold=True, color=BLUE, align=WD_ALIGN_PARAGRAPH.RIGHT, before=3, after=3)


def add_shot_table(doc, scene_index, scene, global_start, shot_no):
    doc.add_page_break()
    end = global_start + scene["duration"]
    add_heading(doc, f"第{scene_index}场  {scene['title']}  |  {timecode(global_start)}-{timecode(end)}  |  {scene['duration']}秒")
    add_paragraph(doc, f"地点：{scene['location']}  ｜  本场事件：{scene['beat']}", size=8, color="4B5563", keep=True, after=3)

    table = doc.add_table(rows=1, cols=7)
    headers = ["镜号", "时间码", "秒", "景别／机位", "画面与动作", "声音／对白", "制作与专业控制"]
    for i, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], h, 7.4, bold=True, color=WHITE, align=WD_ALIGN_PARAGRAPH.CENTER)

    cursor = global_start
    for duration, shot_type, visual, sound, control in scene["shots"]:
        cells = table.add_row().cells
        values = [f"{shot_no:03d}", f"{timecode(cursor)}–{timecode(cursor + duration)}", duration, shot_type, visual, sound, control]
        for cidx, value in enumerate(values):
            set_cell_text(cells[cidx], value, 7.25, bold=False,
                          align=WD_ALIGN_PARAGRAPH.CENTER if cidx in (0, 1, 2) else WD_ALIGN_PARAGRAPH.LEFT)
        cursor += duration
        shot_no += 1

    style_table(table, [0.48, 0.95, 0.38, 1.18, 2.82, 2.55, 2.43], header=True, font_size=7.25)
    return end, shot_no


def add_professional_check(doc):
    doc.add_page_break()
    add_heading(doc, "四、专业性检查与待决事项")
    add_paragraph(doc, "检查口径：公开工程资料优先；未公开部分只做必要、可逆的影视化推理，并在表中显式标注。",
                  size=8.3, color="4B5563", after=4)
    table = doc.add_table(rows=1, cols=3)
    for i, h in enumerate(["检查项", "结论", "复核说明"]):
        set_cell_text(table.rows[0].cells[i], h, 8, bold=True, color=WHITE, align=WD_ALIGN_PARAGRAPH.CENTER)
    for item, result, note in PROFESSIONAL_CHECKS:
        cells = table.add_row().cells
        set_cell_text(cells[0], item, 7.7, bold=True, color=BLUE)
        set_cell_text(cells[1], result, 7.7, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(cells[2], note, 7.7)
        if result == "通过":
            set_cell_shading(cells[1], GREEN)
        elif "待决" in result:
            set_cell_shading(cells[1], RED)
        else:
            set_cell_shading(cells[1], YELLOW)
    style_table(table, [1.45, 1.25, 8.09], header=True, font_size=7.7)

    add_paragraph(doc, "专业检查结论", size=9.5, bold=True, color=BLUE, keep=True, before=5)
    add_paragraph(
        doc,
        "V1.3已把三段直播结构、48小时准备、着陆后6小时整备、约28小时月面驻留、月面起飞再对接、三天地月返回、跳跃式再入和东风搜救统一为一条闭合叙事链。鹊桥二号、天链、深空地面站和黑障期间外测的能力边界均按公开资料重新校正。具体未公开参数继续作为艺术推演，不进入官方事实口径。",
        size=8.5,
        color="26384A",
        after=4,
    )


def add_sources(doc):
    add_heading(doc, "五、公开资料依据与艺术化边界")
    add_paragraph(doc, "下列来源只用于确定公开构型、已验证流程和材料事实；片中日期、人物、呼号、口令、具体参数、回收地点和未公开任务构型不代表官方方案。",
                  size=8.2, color="4B5563", after=3)
    table = doc.add_table(rows=1, cols=2)
    for i, h in enumerate(["资料", "链接"]):
        set_cell_text(table.rows[0].cells[i], h, 8, bold=True, color=WHITE, align=WD_ALIGN_PARAGRAPH.CENTER)
    for title, url in SOURCES:
        cells = table.add_row().cells
        set_cell_text(cells[0], title, 7.5, bold=True, color=BLUE)
        set_cell_text(cells[1], url, 7.1, color="245A8D")
    style_table(table, [4.45, 6.34], header=True, font_size=7.4)

    add_paragraph(doc, "V1.3修订要点", size=9.5, bold=True, color=BLUE, keep=True, before=5)
    revision_points = [
        "1. 全片由连续工程流程重构为三段直播，加入明确的直播进口、直播出口和时间跨越。",
        "2. 增加一名特派记者和飞行任务、乘员系统、月球地质三个专业方向专家。",
        "3. 锁定48小时准备、着陆后6小时出舱整备、约28小时月面总驻留和一次约7小时40分主舱外活动。",
        "4. 把重复的多方确认压缩为关键状态闭合；关键动作期间停止记者和专家覆盖。",
        "5. 修正环月轨道把地球表现为蓝色星点的问题，统一为可分辨的蓝白色圆盘。",
        "6. 主着陆场锁定东风，补齐再入前着陆场展开、黑障外测、终端红外、搜救和医监医保。",
        "7. 鹊桥二号更正为环月大椭圆冻结轨道专用中继星；不采用鹊桥三号。",
        "8. 天链不作为环月或月面主链路，只在地球近域保留可选辅助边界。",
        "9. 删除‘进入地球同步轨道’和‘以第二宇宙速度撞向大气层’等错误说法。",
        "10. 延续全片无配乐、太空和月面外部无空气传播声及后期字幕策略。",
    ]
    for point in revision_points:
        add_paragraph(doc, point, size=8.2, color="26384A", after=1.5)


def build():
    total, total_shots = validate()
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    doc = Document(str(SOURCE))
    clear_body(doc)
    configure_document(doc)
    doc.core_properties.title = "《奔月》正式分镜稿 V1.3"
    doc.core_properties.subject = "中国首次载人登月未来纪实推演"
    doc.core_properties.author = "《奔月》项目组"
    doc.core_properties.comments = "以V1.2版式为模板重构；三段直播结构与公开资料边界见正文。"

    add_cover(doc, total_shots)
    add_control_section(doc)
    add_roles_section(doc)
    add_overview(doc)

    global_start = 0
    shot_no = 1
    for idx, scene in enumerate(SCENES, start=1):
        global_start, shot_no = add_shot_table(doc, idx, scene, global_start, shot_no)

    add_professional_check(doc)
    add_sources(doc)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUTPUT))

    check_doc = Document(str(OUTPUT))
    scene_headings = [p.text for p in check_doc.paragraphs if p.text.startswith("第") and "场" in p.text]
    table_count = len(check_doc.tables)
    all_text = "\n".join(p.text for p in check_doc.paragraphs)
    for table in check_doc.tables:
        for row in table.rows:
            all_text += "\n" + "\t".join(cell.text for cell in row.cells)

    forbidden = [
        "这面国旗，也是一次面向未来月球建设的材料预演。",
        "服务舱观察室",
        "背景音乐",
        "配乐起",
        "音乐渐",
    ]
    found_forbidden = [term for term in forbidden if term in all_text]
    section = check_doc.sections[0]
    qa = {
        "output": str(OUTPUT),
        "bytes": OUTPUT.stat().st_size,
        "target_seconds": TARGET_SECONDS,
        "validated_seconds": total,
        "scene_count": len(SCENES),
        "shot_count": total_shots,
        "next_shot_number": shot_no,
        "scene_heading_count": len(scene_headings),
        "table_count": table_count,
        "section_count": len(check_doc.sections),
        "orientation": str(section.orientation),
        "page_width_in": round(section.page_width.inches, 2),
        "page_height_in": round(section.page_height.inches, 2),
        "forbidden_terms_found": found_forbidden,
        "required_terms": {
            "小天": "小天" in all_text,
            "半弹道跳跃式再入": "半弹道跳跃式再入" in all_text,
            "玄武岩纤维": "玄武岩纤维" in all_text,
            "无配乐": "无配乐" in all_text,
        "2029年4月8日": "2029年4月8日" in all_text,
        "11.97°N": "11.97°N" in all_text,
        "东风着陆场": "东风着陆场" in all_text,
        "鹊桥二号": "鹊桥二号" in all_text,
        "天链": "天链" in all_text,
        "T+28": "T+28" in all_text,
        },
    }
    if global_start != TARGET_SECONDS:
        raise ValueError(f"全局时间码未闭合：{global_start}")
    if shot_no != total_shots + 1:
        raise ValueError(f"镜号未闭合：{shot_no}")
    if len(scene_headings) != len(SCENES):
        raise ValueError(f"场次标题数量错误：{len(scene_headings)}")
    if found_forbidden:
        raise ValueError(f"发现禁用文本：{found_forbidden}")
    if not all(qa["required_terms"].values()):
        raise ValueError(f"必要内容缺失：{qa['required_terms']}")
    QA_OUTPUT.write_text(json.dumps(qa, ensure_ascii=False, indent=2), encoding="utf-8")
    return qa


if __name__ == "__main__":
    report = build()
    print(json.dumps(report, ensure_ascii=False, indent=2))
