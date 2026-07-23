from __future__ import annotations

import json
import re
import sys
import zipfile
from pathlib import Path

from docx import Document


HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
if str(HERE) not in sys.path:
    sys.path.insert(0, str(HERE))

from v131_content import FINAL_TARGET_SECONDS, SCENES, TARGET_SECONDS, validate  # noqa: E402


DOCX = ROOT / "文本" / "奔月_正式分镜稿_V1.3.1.docx"
OUTPUT = HERE / "v131_independent_verification.json"


def collect_text(doc: Document) -> str:
    items = [paragraph.text for paragraph in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            items.extend(cell.text for cell in row.cells)
    return "\n".join(items)


def main() -> dict:
    source_seconds, source_shots = validate()
    doc = Document(str(DOCX))
    all_text = collect_text(doc)

    scene_headings = [
        paragraph.text
        for paragraph in doc.paragraphs
        if re.match(r"^第\d+场\s+", paragraph.text)
    ]
    shot_ids = []
    durations = []
    empty_shot_cells = []
    for table_index, table in enumerate(doc.tables):
        for row_index, row in enumerate(table.rows):
            values = [cell.text.strip() for cell in row.cells]
            if values and re.fullmatch(r"\d{3}", values[0]):
                shot_ids.append(values[0])
                durations.append(int(values[2]))
                if any(not value for value in values):
                    empty_shot_cells.append([table_index, row_index, values[0]])

    with zipfile.ZipFile(DOCX) as package:
        header_footer = {
            name: package.read(name).decode("utf-8", errors="replace")
            for name in package.namelist()
            if name.startswith("word/header") or name.startswith("word/footer")
        }

    required_terms = {
        "表演与剪辑余量版": "表演与剪辑余量版" in all_text,
        "最终剪辑目标14分钟": "最终剪辑目标 14分00秒" in all_text,
        "建议生成素材17分钟": "建议生成素材 17分00秒" in all_text,
        "揽月无人先期入轨": "无人状态的揽月先期升空" in all_text,
        "三人同乘梦舟": "梦舟载三名航天员启程" in all_text,
        "梦舟三人到一人再三人": all(
            phrase in all_text
            for phrase in ("第一次建立三座满员", "两张座椅空置", "完成视觉母题回收")
        ),
        "视觉重建标识": "基于实时遥测的视觉重建" in all_text,
        "真实月面相机": "月面相机固定" in all_text and "同一台已部署相机" in all_text,
        "第一步无虚构口号": "不虚构口号" in all_text,
        "月面时间闭合": all(term in all_text for term in ("T+06:00", "T+13:40", "T+28:00")),
        "阿波罗约六个半小时": "阿波罗11号大约在着陆六个半小时后" in all_text,
        "着陆区理由": "兼顾工程可达性和科学价值" in all_text,
        "返航持续工作": "姿态与热控、生命保障" in all_text,
        "东风搜救时间压缩": "时间压缩｜搜救分队抵达并完成初步安全处置" in all_text,
        "重力再适应": "航天员不能立即独立站立" in all_text,
        "黑障外测": all(term in all_text for term in ("雷达", "光学", "红外发现目标")),
        "无配乐": "全片无配乐" in all_text,
        "月面外部无空气声": "月面外部无空气传播声" in all_text,
    }
    forbidden_terms = {
        "对接前揽月载人": "揽月舱内固定中景，杨凛和陈砚固定在座椅" in all_text,
        "梦舟仅刘一人迎接": "刘定槎从中央工作位观察闭合速度和姿态误差，另外两张座椅空置" in all_text,
        "环月地球小蓝点": "细小蓝点" in all_text,
        "虚构月面风声": "使用旗帜飘动声" in all_text,
        "背景音乐": "背景音乐" in all_text,
        "鹊桥三号启用": "采用鹊桥三号" in all_text,
    }

    expected_ids = [f"{index:03d}" for index in range(1, source_shots + 1)]
    checks = {
        "docx_exists": DOCX.exists(),
        "bytes": DOCX.stat().st_size,
        "source_material_seconds": source_seconds,
        "docx_shot_duration_sum": sum(durations),
        "material_target_seconds": TARGET_SECONDS,
        "final_edit_target_seconds": FINAL_TARGET_SECONDS,
        "source_scene_count": len(SCENES),
        "docx_scene_heading_count": len(scene_headings),
        "source_shot_count": source_shots,
        "docx_shot_count": len(shot_ids),
        "shot_ids_contiguous": shot_ids == expected_ids,
        "shot_ids_unique": len(shot_ids) == len(set(shot_ids)),
        "empty_shot_cells": empty_shot_cells,
        "table_count": len(doc.tables),
        "section_count": len(doc.sections),
        "metadata_title": doc.core_properties.title,
        "required_terms": required_terms,
        "forbidden_terms_found": [name for name, found in forbidden_terms.items() if found],
        "header_footer_v131": all("V1.3.1" in xml for xml in header_footer.values()),
        "header_footer_parts": sorted(header_footer),
    }
    checks["passed"] = all([
        checks["source_material_seconds"] == checks["docx_shot_duration_sum"] == checks["material_target_seconds"] == 1020,
        checks["final_edit_target_seconds"] == 840,
        checks["source_scene_count"] == checks["docx_scene_heading_count"] == 15,
        checks["source_shot_count"] == checks["docx_shot_count"] == 119,
        checks["shot_ids_contiguous"],
        checks["shot_ids_unique"],
        not checks["empty_shot_cells"],
        all(required_terms.values()),
        not checks["forbidden_terms_found"],
        checks["header_footer_v131"],
    ])
    OUTPUT.write_text(json.dumps(checks, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(checks, ensure_ascii=False, indent=2))
    if not checks["passed"]:
        raise SystemExit(1)
    return checks


if __name__ == "__main__":
    main()
