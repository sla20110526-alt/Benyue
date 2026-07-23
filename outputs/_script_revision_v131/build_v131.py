from __future__ import annotations

import json
import re
import sys
from pathlib import Path

from docx import Document


HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
V13_DIR = ROOT / "outputs" / "_script_revision_v13"
for path in (ROOT, HERE, V13_DIR):
    if str(path) not in sys.path:
        sys.path.insert(0, str(path))

import build_v13 as b  # noqa: E402
import v131_content as c  # noqa: E402


SOURCE = ROOT / "文本" / "奔月_正式分镜稿_V1.3.docx"
OUTPUT = ROOT / "文本" / "奔月_正式分镜稿_V1.3.1.docx"
QA_OUTPUT = HERE / "v131_qa_report.json"


for name in (
    "ARCHITECTURE",
    "CREW",
    "GROUND",
    "LANDING_DATE",
    "LANDING_SITE",
    "MEDIA",
    "PROFESSIONAL_CHECKS",
    "RETURN_SITE",
    "SCENES",
    "SOURCES",
    "SUBTITLE",
    "TARGET_SECONDS",
    "TITLE",
    "VERSION",
):
    setattr(b, name, getattr(c, name))


def add_cover(doc: Document, total_shots: int):
    b.add_paragraph(doc, c.TITLE, size=26, bold=True, color=b.BLUE, align=b.WD_ALIGN_PARAGRAPH.CENTER, before=8, after=2)
    b.add_paragraph(doc, c.SUBTITLE, size=12, color=b.MID_BLUE, align=b.WD_ALIGN_PARAGRAPH.CENTER, after=4)
    b.add_paragraph(
        doc,
        f"{c.VERSION}  |  最终剪辑目标 14分00秒  |  建议生成素材 17分00秒  |  {len(c.SCENES)}场  |  {total_shots}镜",
        size=10.5,
        bold=True,
        color="26384A",
        align=b.WD_ALIGN_PARAGRAPH.CENTER,
        after=6,
    )
    b.add_paragraph(
        doc,
        "说明：本版是表演与剪辑余量版。分镜秒数表示建议生成素材长度，包含起落句、呼吸、反应、动作完成和剪辑手柄，不等于最终成片逐镜硬时长；最终剪辑仍以约14分钟为目标。片头和片尾推演声明纳入建议素材时长。",
        size=8.5,
        color="4B5563",
        align=b.WD_ALIGN_PARAGRAPH.CENTER,
        after=5,
    )
    b.add_kv_table(doc, [
        ("正式基线", "文本/奔月_正式分镜稿_V1.3.1.docx；V1.3及以前版本转为历史基线。"),
        ("修订源稿", "文本/奔月_正式分镜稿_V1.3.docx；本次不覆盖原稿。"),
        ("任务日期", c.LANDING_DATE),
        ("着陆地点", c.LANDING_SITE),
        ("返回地点", c.RETURN_SITE),
        ("画面规格", "16:9；成片3840×2160；24fps；未来纪实、真实工程质感、AI真人。"),
        ("声音锁定", "全片无配乐；对白、无线电、设备声和任务通信随镜头同步生成；月面外部无空气传播声。"),
        ("直播结构", "三段直播：月轨对接；落月—月面任务—再对接；离月返回—再入—东风回收。长时间段用字幕、任务回放和可追溯时间压缩跨越。"),
        ("艺术光照", "保留2029年4月8日与锁定坐标，沿用已确认的艺术化月面硬日光，并在片头片尾推演声明中承担偏差。"),
    ])


def add_control_section(doc: Document):
    b.add_heading(doc, "一、全片制作总控")
    b.add_kv_table(doc, [
        ("核心叙事", "三名航天员同乘梦舟抵达月轨，两人转入揽月落月，一人留守梦舟；两人携样品返回、三人重聚并平安返航。"),
        ("五个戏剧问题", "两器能否会合；两人能否安全着陆；中国航天员的第一步如何发生；他们在月面做什么并能否返回梦舟；三人能否穿越黑障并被安全回收。"),
        ("座椅视觉母题", "梦舟三座满员—两张空座—三座重聚—三座返航；用于交代乘组关系，不用额外对白重复说明。"),
        ("表演余量", "采访回答优先保留10—16秒；任务口令保留反应与通信延迟；关键动作完成后再切镜。最终剪辑可压缩，但生成时不以14分钟硬卡台词。"),
        ("时间压缩", "48小时准备、着陆后6小时整备、约28小时月面驻留、起飞后约3小时40分再交会和约三天地月返回均使用明确字幕、任务相机时间码和匹配剪辑。"),
        ("AI边界", "小天负责传感器融合、走廊预测、候选区排序和异常提示；最终选择、授权和中止权归乘组与北京。"),
        ("声学边界", "全片无配乐；太空和月面外部不使用空气传播的发动机声、脚步声、尘土声、车轮声或旗帜飘动声。"),
        ("文字策略", "任务日期、地点、时间码、视觉重建标识和推演声明由后期字幕完成；AI原始画面不生成可读文字。"),
        ("直播解释", "记者和专家只在事件前后解释；对接、最终下降、第一步、月面起飞、再对接、黑障恢复和着陆由任务通信主导。"),
    ])
    b.add_paragraph(doc, "1.1 三类视角及使用边界", size=9, bold=True, color=b.BLUE, keep=True, before=3)
    table = doc.add_table(rows=1, cols=2)
    for i, heading in enumerate(["视角", "画面来源与叙事职责"]):
        b.set_cell_text(table.rows[0].cells[i], heading, 8, bold=True, color=b.WHITE, align=b.WD_ALIGN_PARAGRAPH.CENTER)
    for name, description in c.VIEWPOINT_RULES:
        cells = table.add_row().cells
        b.set_cell_text(cells[0], name, 7.8, bold=True, color=b.BLUE)
        b.set_cell_text(cells[1], description, 7.8)
    b.style_table(table, [1.65, 9.14], header=True, font_size=7.8)


def add_overview(doc: Document):
    b.add_heading(doc, "三、场次与建议素材时长总览")
    table = doc.add_table(rows=1, cols=6)
    headers = ["场次", "场名", "建议素材时间码", "素材秒数", "地点", "主要事件点"]
    for i, heading in enumerate(headers):
        b.set_cell_text(table.rows[0].cells[i], heading, 8, bold=True, color=b.WHITE, align=b.WD_ALIGN_PARAGRAPH.CENTER)
    cursor = 0
    for index, scene in enumerate(c.SCENES, start=1):
        cells = table.add_row().cells
        values = [
            index,
            scene["title"],
            f"{b.timecode(cursor)}–{b.timecode(cursor + scene['duration'])}",
            f"{scene['duration']}秒",
            scene["location"],
            scene["beat"],
        ]
        for column, value in enumerate(values):
            b.set_cell_text(
                cells[column],
                value,
                7.4,
                bold=(column == 1),
                color=(b.BLUE if column == 1 else "1A1A1A"),
                align=b.WD_ALIGN_PARAGRAPH.CENTER if column in (0, 2, 3) else b.WD_ALIGN_PARAGRAPH.LEFT,
            )
        cursor += scene["duration"]
    b.style_table(table, [0.48, 1.18, 1.12, 0.58, 2.08, 5.35], header=True, font_size=7.4)
    b.add_paragraph(
        doc,
        f"建议生成素材：{len(c.SCENES)}场，{sum(len(scene['shots']) for scene in c.SCENES)}镜，"
        f"{c.TARGET_SECONDS}秒（17分00秒）；最终剪辑目标约{c.FINAL_TARGET_SECONDS}秒（14分00秒）。",
        size=8.5,
        bold=True,
        color=b.BLUE,
        align=b.WD_ALIGN_PARAGRAPH.RIGHT,
        before=3,
        after=3,
    )


def add_professional_check(doc: Document):
    doc.add_page_break()
    b.add_heading(doc, "四、专业性检查与待决事项")
    b.add_paragraph(
        doc,
        "检查口径：公开工程资料优先；未公开部分只做必要、可逆的影视化推理。不可逆事件由任务证据机位确认；视觉重建必须显式标注。",
        size=8.3,
        color="4B5563",
        after=4,
    )
    table = doc.add_table(rows=1, cols=3)
    for i, heading in enumerate(["检查项", "结论", "复核说明"]):
        b.set_cell_text(table.rows[0].cells[i], heading, 8, bold=True, color=b.WHITE, align=b.WD_ALIGN_PARAGRAPH.CENTER)
    for item, result, note in c.PROFESSIONAL_CHECKS:
        cells = table.add_row().cells
        b.set_cell_text(cells[0], item, 7.7, bold=True, color=b.BLUE)
        b.set_cell_text(cells[1], result, 7.7, bold=True, align=b.WD_ALIGN_PARAGRAPH.CENTER)
        b.set_cell_text(cells[2], note, 7.7)
        if result == "通过":
            b.set_cell_shading(cells[1], b.GREEN)
        elif "待决" in result:
            b.set_cell_shading(cells[1], b.RED)
        else:
            b.set_cell_shading(cells[1], b.YELLOW)
    b.style_table(table, [1.45, 1.25, 8.09], header=True, font_size=7.7)
    b.add_paragraph(doc, "专业检查结论", size=9.5, bold=True, color=b.BLUE, keep=True, before=5)
    b.add_paragraph(
        doc,
        "V1.3.1修正了对接前乘组位置矛盾，重建了梦舟三座视觉连续性；把直播解释、任务证据和第三视角分成可追溯的三套画面语言；闭合T+06:00出舱、T+13:40结束主舱外活动和T+28:00月面起飞；补足着陆区选择、地月返航工作、黑障外测接力、搜救时间压缩和返回后医学转移。本版秒数是表演与剪辑余量，不改变最终约14分钟的成片目标。",
        size=8.5,
        color="26384A",
        after=4,
    )


def add_sources(doc: Document):
    b.add_heading(doc, "五、公开资料依据与艺术化边界")
    b.add_paragraph(
        doc,
        "下列来源只用于确定公开构型、已验证流程和材料事实；片中日期、人物、呼号、具体参数、着陆坐标和未公开任务构型不代表官方方案。",
        size=8.2,
        color="4B5563",
        after=3,
    )
    table = doc.add_table(rows=1, cols=2)
    for i, heading in enumerate(["资料", "链接"]):
        b.set_cell_text(table.rows[0].cells[i], heading, 8, bold=True, color=b.WHITE, align=b.WD_ALIGN_PARAGRAPH.CENTER)
    for title, url in c.SOURCES:
        cells = table.add_row().cells
        b.set_cell_text(cells[0], title, 7.5, bold=True, color=b.BLUE)
        b.set_cell_text(cells[1], url, 7.1, color="245A8D")
    b.style_table(table, [4.45, 6.34], header=True, font_size=7.4)

    b.add_paragraph(doc, "V1.3.1修订要点", size=9.5, bold=True, color=b.BLUE, keep=True, before=5)
    revision_points = [
        "1. 修正揽月无人先期入轨与三人乘组位置矛盾：三人同乘梦舟，对接后两人转入揽月。",
        "2. 建立梦舟三座满员—两座空置—三座重聚—三座返航的视觉母题。",
        "3. 全片由14分钟硬时长改为17分钟建议生成素材；最终剪辑目标仍约14分钟。",
        "4. 记者与专家采访增加起落句、停顿、反应和画外跨接，避免专业信息抢读。",
        "5. 月面相机部署前只使用显式标注的遥测视觉重建；部署后使用同一便携月面相机。",
        "6. 第一脚仍由揽月外部固定相机和宇航服相机证明；第三视角不替代关键证据。",
        "7. 以同一台月面相机记录国旗、科研时间压缩和T+28小时起飞，并保留起飞后的月面空镜。",
        "8. 闭合T+06:00首次出舱、T+13:40结束主舱外活动和T+28:00月面起飞。",
        "9. 补充着陆区为何兼顾工程可达性和科学价值；项目坐标仍属于艺术锁定点。",
        "10. 修正无人任务“只需自检”和三天地月返回“相对平静”的过度简化。",
        "11. 黑障期用乘组承压、飞控预测、雷达／光学搜索和终端红外捕获补齐叙事。",
        "12. 搜救抵达以字幕压缩时间；开舱后按重力再适应评估使用支撑座椅或担架转移。",
        "13. 延续全片无配乐、太空与月面外部无空气传播声，以及后期字幕策略。",
    ]
    for point in revision_points:
        b.add_paragraph(doc, point, size=8.2, color="26384A", after=1.5)


def collect_text(doc: Document) -> str:
    items = [paragraph.text for paragraph in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            items.extend(cell.text for cell in row.cells)
    return "\n".join(items)


def build() -> dict:
    total, total_shots = c.validate()
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)

    doc = Document(str(SOURCE))
    b.clear_body(doc)
    b.configure_document(doc)
    doc.core_properties.title = "《奔月》正式分镜稿 V1.3.1"
    doc.core_properties.subject = "中国首次载人登月未来纪实推演·表演与剪辑余量版"
    doc.core_properties.author = "《奔月》项目组"
    doc.core_properties.comments = "以V1.3版式为模板优化；17分钟为建议生成素材，最终剪辑目标约14分钟。"

    add_cover(doc, total_shots)
    add_control_section(doc)
    b.add_roles_section(doc)
    add_overview(doc)

    global_start = 0
    shot_no = 1
    for index, scene in enumerate(c.SCENES, start=1):
        global_start, shot_no = b.add_shot_table(doc, index, scene, global_start, shot_no)

    add_professional_check(doc)
    add_sources(doc)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUTPUT))

    check_doc = Document(str(OUTPUT))
    all_text = collect_text(check_doc)
    scene_headings = [
        paragraph.text
        for paragraph in check_doc.paragraphs
        if re.match(r"^第\d+场\s+", paragraph.text)
    ]
    shot_ids = []
    empty_shot_cells = []
    for table_index, table in enumerate(check_doc.tables):
        for row_index, row in enumerate(table.rows):
            values = [cell.text.strip() for cell in row.cells]
            if values and re.fullmatch(r"\d{3}", values[0]):
                shot_ids.append(values[0])
                if any(not value for value in values):
                    empty_shot_cells.append([table_index, row_index, values[0]])

    required_terms = {
        "V1.3.1": "V1.3.1" in all_text,
        "建议生成素材": "建议生成素材" in all_text,
        "揽月无人": "揽月无人" in all_text,
        "三把座椅": "三把座椅" in all_text,
        "基于实时遥测的视觉重建": "基于实时遥测的视觉重建" in all_text,
        "T+13:40": "T+13:40" in all_text,
        "T+28:00": "T+28:00" in all_text,
        "东风着陆场": "东风着陆场" in all_text,
        "半弹道跳跃式再入": "半弹道跳跃式再入" in all_text,
        "无配乐": "无配乐" in all_text,
    }
    forbidden_terms = {
        "揽月舱内固定中景且杨凛陈砚已在舱内": "揽月舱内固定中景，杨凛和陈砚固定在座椅" in all_text,
        "环月地球小蓝点": "细小蓝点" in all_text,
        "无人探测器只需落地自检": "无人探测器只需落地自检" in all_text,
        "三天地月返回相对平静": "三天地月返回相对平静" in all_text,
        "背景音乐": "背景音乐" in all_text,
    }

    qa = {
        "output": str(OUTPUT),
        "bytes": OUTPUT.stat().st_size,
        "material_target_seconds": c.TARGET_SECONDS,
        "validated_seconds": total,
        "final_edit_target_seconds": c.FINAL_TARGET_SECONDS,
        "scene_count": len(c.SCENES),
        "shot_count": total_shots,
        "docx_scene_heading_count": len(scene_headings),
        "docx_shot_count": len(shot_ids),
        "shot_ids_contiguous": shot_ids == [f"{index:03d}" for index in range(1, total_shots + 1)],
        "empty_shot_cells": empty_shot_cells,
        "table_count": len(check_doc.tables),
        "section_count": len(check_doc.sections),
        "required_terms": required_terms,
        "forbidden_terms_found": [name for name, found in forbidden_terms.items() if found],
    }
    qa["passed"] = all([
        total == c.TARGET_SECONDS == global_start == 1020,
        c.FINAL_TARGET_SECONDS == 840,
        len(c.SCENES) == len(scene_headings) == 15,
        total_shots == len(shot_ids) == 119,
        shot_no == total_shots + 1,
        qa["shot_ids_contiguous"],
        not empty_shot_cells,
        all(required_terms.values()),
        not qa["forbidden_terms_found"],
    ])
    if not qa["passed"]:
        raise ValueError(json.dumps(qa, ensure_ascii=False, indent=2))
    QA_OUTPUT.write_text(json.dumps(qa, ensure_ascii=False, indent=2), encoding="utf-8")
    return qa


if __name__ == "__main__":
    print(json.dumps(build(), ensure_ascii=False, indent=2))
