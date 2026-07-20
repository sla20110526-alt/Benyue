from __future__ import annotations

import json
import re
import zipfile
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[2]
DOCX = ROOT / "文本" / "奔月_正式分镜稿_V1.2.docx"
REPORT = Path(__file__).with_name("v12_independent_verification.json")


def seconds(tc: str) -> int:
    minute, second = tc.split(":")
    return int(minute) * 60 + int(second)


def main():
    with zipfile.ZipFile(DOCX) as archive:
        corrupt_member = archive.testzip()
        package_files = archive.namelist()

    doc = Document(str(DOCX))
    rows = []
    empty_cells = []
    for table_index, table in enumerate(doc.tables, start=1):
        if not table.rows:
            continue
        header = [cell.text.strip() for cell in table.rows[0].cells]
        if header[:3] != ["镜号", "时间码", "秒"]:
            continue
        for row_index, row in enumerate(table.rows[1:], start=2):
            values = [cell.text.strip() for cell in row.cells]
            for cell_index, value in enumerate(values, start=1):
                if not value:
                    empty_cells.append([table_index, row_index, cell_index])
            rows.append(values)

    errors = []
    previous_end = 0
    for expected, values in enumerate(rows, start=1):
        number, tc, duration = values[:3]
        if int(number) != expected:
            errors.append(f"镜号：期望{expected}，实际{number}")
        match = re.fullmatch(r"(\d{2}:\d{2})–(\d{2}:\d{2})", tc)
        if not match:
            errors.append(f"时间码格式错误：{number} {tc}")
            continue
        start, end = map(seconds, match.groups())
        if start != previous_end:
            errors.append(f"时间码不连续：{number} 从{start}开始，上镜结束{previous_end}")
        if end - start != int(duration):
            errors.append(f"镜头时长不符：{number}")
        previous_end = end

    all_text = "\n".join(p.text for p in doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            all_text += "\n" + "\t".join(cell.text for cell in row.cells)
    footer_text = "\n".join(p.text for section in doc.sections for p in section.footer.paragraphs)

    report = {
        "docx": str(DOCX),
        "zip_corrupt_member": corrupt_member,
        "package_file_count": len(package_files),
        "shot_rows": len(rows),
        "final_time_seconds": previous_end,
        "empty_shot_cells": empty_cells,
        "sequence_errors": errors,
        "legacy_version_in_body": "V1.1" in all_text,
        "legacy_version_in_footer": "V1.1" in footer_text,
        "forbidden_exact_sentence_present": "这面国旗，也是一次面向未来月球建设的材料预演。" in all_text,
        "music_cues_present": any(term in all_text for term in ["背景音乐", "配乐起", "音乐渐强", "音乐渐弱"]),
        "scene_heading_count": sum(1 for p in doc.paragraphs if re.match(r"第\d+场", p.text)),
        "status": "PASS",
    }
    if corrupt_member or len(rows) != 114 or previous_end != 840 or empty_cells or errors:
        report["status"] = "FAIL"
    if report["forbidden_exact_sentence_present"] or report["music_cues_present"]:
        report["status"] = "FAIL"
    REPORT.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    if report["status"] != "PASS":
        raise SystemExit(1)


if __name__ == "__main__":
    main()
